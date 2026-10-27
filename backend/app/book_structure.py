"""Extract chapter structure, bibliographic metadata and extraction evidence.

Every format is read by the extractor that understands it: DTBook and EPUB by
their semantic markup, DOCX by its body (tables included), text PDFs by their
layout, and pages that carry no text layer by OCR. Whatever an extractor leaves
out on purpose (running heads, page numbers, note markers, navigation) is
recorded with a reason in :class:`app.extraction_qa.ExtractionReport`, and the
report counts the source's words independently, so content that disappears
without a reason is reported instead of going unnoticed.

Chapter titles an extractor has to invent ("Página 3", "Capítulo 2", the
"Prólogo" before a plain-text book's first chapter) are marked
``title_generated``: they label navigation but are never narrated, because the
book does not say them.
"""

from __future__ import annotations

import io
import re
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

import ebooklib
import fitz
from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from ebooklib import epub

from app import extraction_qa, ocr
from app.extraction_qa import BLOCKING, INFO, WARNING, ExtractionReport


@dataclass
class BookChapter:
    title: str
    paragraphs: list[str] = field(default_factory=list)
    # True when the extractor invented the title; it is navigation, not text.
    title_generated: bool = False
    # Where each paragraph came from, parallel to ``paragraphs`` when known:
    # {"page": 12, "element": "p", "method": "pdf-layout"}. A heading inside
    # the chapter carries "element": "heading" and its "heading_level"; a
    # paragraph at which a page of the original begins lists that page in
    # "page_breaks".
    sources: list[dict] = field(default_factory=list)
    # The same, for the chapter's own title.
    title_source: dict = field(default_factory=dict)

    def source_of(self, index: int) -> dict:
        return self.sources[index] if len(self.sources) == len(self.paragraphs) else {}


@dataclass
class BookStructure:
    title: str
    author: str = ""
    edition: str = ""
    publication_year: str = ""
    publisher: str = ""
    chapters: list[BookChapter] = field(default_factory=list)
    extraction: ExtractionReport | None = None

    def as_text(self) -> str:
        blocks: list[str] = []
        for chapter in self.chapters:
            if chapter.title:
                blocks.append(chapter.title)
            blocks.extend(chapter.paragraphs)
        return "\n\n".join(blocks)

    @property
    def generated_titles(self) -> set[str]:
        return {chapter.title for chapter in self.chapters if chapter.title_generated}


def extract_book_structure(
    filename: str, content: bytes, fallback_text: str = "", *, pages: str | None = None,
) -> BookStructure:
    """Preserve chapter boundaries where the source format exposes them.

    ``pages`` ("20-114") limits a PDF to that page range, as numbered by a PDF
    viewer; other formats have no pages and ignore it.
    """
    path = Path(filename)
    suffix = path.suffix.lower()
    title = path.stem or "Libro"
    result = BookStructure(title=title)
    report = ExtractionReport(source_format=suffix.lstrip(".") or "desconocido")
    result.extraction = report
    raw: list[str] = []

    if suffix == ".xml":
        _from_xml(content, result, fallback_text, report, raw)
    elif suffix == ".epub":
        _from_epub(filename, content, result, fallback_text, report, raw)
    elif suffix == ".docx":
        _from_docx(content, result, fallback_text, report, raw)
    elif suffix == ".pdf":
        _from_pdf(content, result, fallback_text, pages, report, raw)
    elif suffix == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("El archivo de texto no usa una codificacion compatible.") from exc
        report.method = "txt"
        raw.append(text)
        result.chapters = _split_text_chapters(result.title, _drop_transcriber_markers(text, report))
    else:
        result.chapters = [BookChapter(result.title, _paragraphs(fallback_text), title_generated=True)]
        raw.append(fallback_text)
    _finalize(result, report, raw)
    return result


def _finalize(result: BookStructure, report: ExtractionReport, raw: list[str]) -> None:
    extracted: list[str] = []
    for chapter in result.chapters:
        if chapter.title and not chapter.title_generated:
            extracted.append(chapter.title)
        extracted.extend(chapter.paragraphs)
        if chapter.title_generated and chapter.title:
            report.generated_titles.append(chapter.title)
    if report.generated_titles:
        report.issue(INFO, "generated_titles",
                     f"{len(report.generated_titles)} títulos de navegación no están en el libro; "
                     "sirven para navegar y no se narran.", text=", ".join(report.generated_titles[:10]))
    extraction_qa.measure_coverage(report, raw, extracted)
    extraction_qa.inspect_text_quality(report, [(chapter.title, chapter.paragraphs) for chapter in result.chapters])


# --------------------------------------------------------------------------- #
# Block flattening: every text-bearing block, once, in reading order
# --------------------------------------------------------------------------- #

class _Flattener:
    """Collects the text of a document as a sequence of blocks.

    Text inside a block, including its inline markup, becomes one paragraph;
    a nested block ends the current paragraph and starts its own, so mixed
    content ("text <list>…</list> more text") loses nothing. Notes are taken
    out of the flow and read at the end of their chapter.
    """

    def __init__(self, report: ExtractionReport, location: str, page: str = "",
                 pending_pages: list[str] | None = None) -> None:
        self.report = report
        self.location = location
        self.page = page
        self.blocks: list[tuple[str, str, dict]] = []  # (kind, text, provenance)
        self.notes: list[str] = []
        # Pages of the original that began since the last block: they begin at
        # the next block, which is where a page marker belongs.
        self.pending_pages: list[str] = list(pending_pages or [])
        # How deep the enclosing <levelN> is, for a <hd> that says no number.
        self.level_depth = 0
        self._current: list[str] = []
        self._open: list[str] = []
        self._levels: list[int | None] = []

    def text(self, value: str | None) -> None:
        if value:
            self._current.append(value)

    def open(self, kind: str, level: int | None = None) -> None:
        self.flush()
        self._open.append(kind)
        self._levels.append(level)

    def close(self) -> None:
        self.flush()
        if self._open:
            self._open.pop()
            self._levels.pop()

    def page_break(self, label: str) -> None:
        """A page of the original begins here (a DTBook pagenum, an EPUB pagebreak)."""
        if label:
            self.page = label
            self.pending_pages.append(label)

    def flush(self) -> None:
        text = _spoken("".join(self._current))
        self._current = []
        if text:
            kind = self._open[-1] if self._open else "p"
            provenance: dict = {"element": kind}
            level = self._levels[-1] if self._levels else None
            if kind == "heading" and level:
                provenance["heading_level"] = level
            if self.page:
                provenance["page"] = self.page
            if self.pending_pages:
                provenance["page_breaks"] = self.pending_pages
                self.pending_pages = []
            self.blocks.append((kind, text, provenance))

    def note(self, text: str) -> None:
        value = note_text(text)
        if value:
            self.notes.append(value)


def note_text(text: str) -> str:
    """A note as it is read: "Nota 1. …", from the note's own number.

    A note that already says "Nota 1." keeps it: prefixing it again read
    "Nota. Nota 1." aloud.
    """
    text = _spoken(text)
    if not text:
        return ""
    if re.match(r"^nota\s+\d{1,3}\s*[.:]", text, flags=re.I):
        return text
    match = re.match(r"^\s*\[?(\d{1,3})\]?[.)\s]+(.*)$", text, flags=re.S)
    return f"Nota {match.group(1)}. {match.group(2).strip()}" if match else f"Nota. {text}"


_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6", "hd", "levelhd", "bridgehead"}
_XML_BLOCKS = _HEADINGS | {
    "p", "para", "paragraph", "li", "item", "line", "dt", "dd", "tr", "caption", "blockquote", "sidebar",
    "div", "poem", "linegroup", "list", "dl", "table", "prodnote", "epigraph", "byline", "dateline",
    "address", "author", "imggroup", "level", "level1", "level2", "level3", "level4", "level5", "level6",
    "section", "chapter", "title", "frontmatter", "bodymatter", "rearmatter", "book", "covertitle",
}
_XML_CELLS = {"td", "th"}
_XML_NOTES = {"note", "annotation", "footnote", "endnote"}
_XML_MARKERS = {"noteref", "annoref"}
_XML_SILENT = {"head", "script", "style", "meta", "link"}


def _walk_xml(element: ET.Element, flat: _Flattener) -> None:
    name = _local_name(element.tag).lower()
    if name in _XML_SILENT:
        flat.text(element.tail)
        return
    if name == "pagenum":
        flat.page_break(_spoken("".join(element.itertext())))
        flat.report.drop("page_number", "".join(element.itertext()), "Número de página del original: marca la "
                         "posición, no se narra.", f"página {flat.page}")
        flat.text(element.tail)
        return
    if name in _XML_MARKERS:
        flat.report.drop("note_marker", "".join(element.itertext()), "Marca de llamada a nota; la nota se lee "
                         "al final del capítulo.", flat.location)
        flat.text(element.tail)
        return
    if name in _XML_NOTES:
        flat.note("".join(element.itertext()))
        flat.text(element.tail)
        return
    if name == "img":
        alt = element.get("alt", "").strip()
        if alt:
            flat.report.drop("image_alt", alt, "Texto alternativo de una imagen; no se narra.", flat.location)
        flat.text(element.tail)
        return
    block = name in _XML_BLOCKS
    depth = int(name[5]) if re.fullmatch(r"level[1-6]", name) else None
    if block:
        flat.open("heading" if name in _HEADINGS or name == "title" else name, _heading_level(name, flat.level_depth))
    outer_depth = flat.level_depth
    if depth:
        flat.level_depth = depth
    flat.text(element.text)
    for child in element:
        _walk_xml(child, flat)
        if _local_name(child.tag).lower() in _XML_CELLS:
            flat.text("; ")
    flat.level_depth = outer_depth
    if block:
        flat.close()
    flat.text(element.tail)


def _heading_level(name: str, level_depth: int) -> int | None:
    """The outline level a heading element states: h2 → 2, a DTBook hd → its level's depth."""
    if re.fullmatch(r"h[1-6]", name):
        return int(name[1])
    if name == "hd" and level_depth:
        return level_depth
    return None


def _chapter_from_blocks(flat: _Flattener, fallback_title: str, method: str) -> BookChapter:
    blocks = list(flat.blocks)
    title, generated = fallback_title, True
    title_source: dict = {}
    if blocks and blocks[0][0] == "heading":
        title, generated = blocks[0][1], False
        title_source = {**blocks[0][2], "method": method}
        blocks = blocks[1:]
    paragraphs = [text for _kind, text, _source in blocks]
    sources = [{**source, "method": method} for _kind, _text, source in blocks]
    paragraphs += flat.notes
    sources += [{"element": "note", "method": method} for _ in flat.notes]
    return BookChapter(title, paragraphs, title_generated=generated, sources=sources, title_source=title_source)


# --------------------------------------------------------------------------- #
# DTBook and generic XML
# --------------------------------------------------------------------------- #

def _from_xml(content: bytes, result: BookStructure, fallback_text: str, report: ExtractionReport,
              raw: list[str]) -> None:
    try:
        root = ET.fromstring(content)
    except ET.ParseError as exc:
        raise ValueError("El archivo XML no está bien formado y no se puede leer con seguridad.") from exc
    metadata: dict[str, str] = {}
    for node in root.iter():
        local = _local_name(node.tag).lower()
        name = (node.get("name") or "").strip().lower()
        value = (node.get("content") or _spoken("".join(node.itertext()))).strip()
        if name and value:
            metadata[name.removeprefix("dc:").removeprefix("ncc:")] = value
        elif local in {"title", "creator", "author", "edition", "publisher", "date"} and value:
            metadata.setdefault(local, value)

    result.title = metadata.get("title", result.title)
    result.author = metadata.get("creator", metadata.get("author", ""))
    result.edition = metadata.get("sourceedition", metadata.get("edition", ""))
    result.publication_year = metadata.get("sourcedate", metadata.get("date", ""))
    result.publisher = metadata.get("sourcepublisher", metadata.get("publisher", ""))

    parents = {child: parent for parent in root.iter() for child in parent}
    nodes = list(root.iter())
    dtbook_levels = [node for node in nodes if _local_name(node.tag).lower() == "level1"]
    is_dtbook = bool(dtbook_levels) or _local_name(root.tag).lower() == "dtbook"
    report.method = "dtbook" if is_dtbook else "xml"
    candidates = dtbook_levels or [node for node in nodes if _local_name(node.tag).lower() == "chapter"]
    top_level = [node for node in candidates if not any(parent in candidates for parent in _ancestors(node, parents))]

    book = next((node for node in nodes if _local_name(node.tag).lower() == "book"), None) if is_dtbook else None
    content_root = book if book is not None else root
    raw.append(" ".join(content_root.itertext()))
    if not top_level:
        top_level = [content_root]
    else:
        # Text outside every chapter is either metadata (doctitle, docauthor,
        # the generic format's <title>) or content a chapter list would lose.
        inside = {id(node) for chapter in top_level for node in chapter.iter()}
        for node in content_root.iter():
            if id(node) in inside or node is content_root:
                continue
            own = _spoken((node.text or "") if len(node) else "".join(node.itertext()))
            if not own:
                continue
            local = _local_name(node.tag).lower()
            if local in {"doctitle", "docauthor", "title", "creator", "author", "edition", "date",
                         "publisher", "covertitle"} or node.get("name"):
                report.drop("metadata", own, "Metadatos del libro, usados en los datos bibliográficos.",
                            local)
            else:
                report.drop("outside_chapter", own, "Texto fuera de cualquier capítulo.", local)
                report.issue(WARNING, "outside_chapter",
                             "Hay texto fuera de los capítulos del libro que no se narrará.", text=own)

    page = ""
    pending: list[str] = []
    for index, chapter in enumerate(top_level, 1):
        flat = _Flattener(report, f"capítulo {index}", page, pending)
        depth = _local_name(chapter.tag).lower()
        flat.level_depth = int(depth[5]) if re.fullmatch(r"level[1-6]", depth) else 1
        flat.open("chapter")
        flat.text(chapter.text)
        for child in chapter:
            _walk_xml(child, flat)
        flat.close()
        page = flat.page
        # A page that begins after the chapter's last block begins with the next chapter.
        pending = flat.pending_pages
        result.chapters.append(_chapter_from_blocks(flat, f"Capítulo {index}", report.method))

    if not result.chapters or not any(chapter.paragraphs for chapter in result.chapters):
        result.chapters = [BookChapter(result.title, _paragraphs(fallback_text), title_generated=True)]


# --------------------------------------------------------------------------- #
# EPUB
# --------------------------------------------------------------------------- #

_NOTE_TYPES = {"footnote", "endnote", "rearnote", "note"}
_HTML_BLOCKS = {
    "p", "div", "li", "ul", "ol", "dl", "dt", "dd", "h1", "h2", "h3", "h4", "h5", "h6", "blockquote", "pre",
    "figcaption", "figure", "caption", "table", "tr", "thead", "tbody", "tfoot", "section", "article", "aside",
    "header", "main", "body", "html", "hgroup", "address", "center", "details", "summary", "legend", "fieldset",
}
_HTML_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_HTML_SILENT = {"script", "style", "head", "title", "meta", "link", "svg", "math"}
_COVER_NAMES = {"cover", "coverpage", "cover-page", "titlepage", "title-page", "portada", "cubierta",
                "couverture", "capa", "frontcover"}


def _epub_semantics(node) -> set[str]:
    values = f"{node.get('epub:type', '')} {node.get('role', '')}".replace("doc-", " ")
    return set(values.split())


def _is_note_marker(node: Tag) -> bool:
    if "noteref" in _epub_semantics(node):
        return True
    # A plain link whose text is only a number pointing inside the book: the
    # reference mark of a note, read aloud as "uno" at the end of a sentence.
    if node.name in {"a", "sup"}:
        label = node.get_text("").strip()
        link = node if node.name == "a" else node.find("a")
        href = link.get("href", "") if link is not None else ""
        return bool(re.fullmatch(r"[\[(]?\d{1,3}[\])]?", label)) and "#" in href
    return False


def _walk_html(node: Tag, flat: _Flattener, title_heading: list) -> None:
    for child in list(node.children):
        if isinstance(child, Comment):
            continue
        if isinstance(child, NavigableString):
            if type(child) is NavigableString:
                flat.text(str(child))
            continue
        if not isinstance(child, Tag):
            continue
        name = (child.name or "").lower()
        semantics = _epub_semantics(child)
        if name in _HTML_SILENT:
            continue
        if name == "nav" or "toc" in semantics:
            flat.report.drop("navigation", child.get_text(" "), "Índice de navegación del EPUB; la navegación "
                             "se genera a partir de los capítulos.", flat.location)
            continue
        if "pagebreak" in semantics:
            flat.page_break(child.get("title", "") or child.get_text("").strip() or child.get("aria-label", ""))
            flat.report.drop("page_number", child.get_text(" ") or child.get("title", ""),
                             "Marca de página del original; no se narra.", flat.location)
            continue
        if _is_note_marker(child):
            flat.report.drop("note_marker", child.get_text(""), "Marca de llamada a nota; la nota se lee al "
                             "final del capítulo.", flat.location)
            continue
        if semantics & _NOTE_TYPES:
            flat.note(_flatten_inline(child))
            continue
        if name == "footer":
            text = child.get_text(" ")
            flat.report.drop("footer", text, "Pie de documento del EPUB; no se narra.", flat.location)
            if text.strip():
                flat.report.issue(WARNING, "epub_footer", "Un pie de documento del EPUB no se narrará; "
                                  "comprueba que no sea contenido del libro.", location=flat.location, text=text)
            continue
        if name == "br":
            flat.text(" ")
            continue
        if name == "img":
            alt = (child.get("alt") or "").strip()
            if alt:
                flat.report.drop("image_alt", alt, "Texto alternativo de una imagen; no se narra.", flat.location)
            continue
        block = name in _HTML_BLOCKS
        if block:
            flat.open("heading" if name in _HTML_HEADINGS else name,
                      int(name[1]) if name in _HTML_HEADINGS else None)
            if name in _HTML_HEADINGS and not title_heading and not flat.blocks:
                title_heading.append(name)
        _walk_html(child, flat, title_heading)
        if name in {"td", "th"}:
            flat.text("; ")
        if block:
            flat.close()


def _flatten_inline(node: Tag) -> str:
    # The note's own number stays: it is what "Nota 1." is read from.
    for line_break in node.find_all("br"):
        line_break.replace_with(" ")
    return node.get_text("")


def _epub_toc_titles(book) -> dict[str, str]:
    """Spine document name -> its title in the book's own table of contents."""
    titles: dict[str, str] = {}

    def visit(entries) -> None:
        for entry in entries or []:
            if isinstance(entry, (tuple, list)) and entry:
                visit([entry[0]])
                if len(entry) > 1:
                    visit(entry[1])
                continue
            href = getattr(entry, "href", "") or ""
            label = getattr(entry, "title", "") or ""
            name = href.split("#", 1)[0]
            if name and label and name not in titles:
                titles[name] = " ".join(label.split())

    try:
        visit(getattr(book, "toc", []) or [])
    except Exception:
        return titles
    return titles


def _from_epub(filename: str, content: bytes, result: BookStructure, fallback_text: str,
               report: ExtractionReport, raw: list[str]) -> None:
    try:
        with tempfile.TemporaryDirectory(prefix="soni2-structure-") as temp_dir:
            source = Path(temp_dir) / (Path(filename).name or "book.epub")
            source.write_bytes(content)
            book = epub.read_epub(str(source))
    except Exception as exc:
        raise ValueError("No hemos podido abrir este EPUB. Comprueba que no este protegido con DRM y que sea un EPUB valido.") from exc

    report.method = "epub"
    result.title = _epub_metadata(book, "title") or result.title
    result.author = _epub_metadata(book, "creator")
    result.edition = _epub_metadata(book, "sourceedition")
    result.publication_year = _epub_metadata(book, "date")
    result.publisher = _epub_metadata(book, "publisher")
    toc_titles = _epub_toc_titles(book)
    page = ""
    pending: list[str] = []
    for index, (item_id, _linear) in enumerate(book.spine, 1):
        item = book.get_item_with_id(item_id)
        if item is None or item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        soup = BeautifulSoup(item.get_content(), "html.parser")
        root = soup.body or soup
        raw_text = " ".join(text for text in root.find_all(string=True)
                            if not isinstance(text, Comment)
                            and not any(parent.name in _HTML_SILENT for parent in text.parents if parent.name))
        item_name = item.get_name()
        item_stem = Path(item_name).stem.casefold()
        if item_stem in _COVER_NAMES:
            report.drop("cover", raw_text, "Cubierta o portada del EPUB; no se narra.", item_name)
            raw.append(raw_text)
            continue
        raw.append(raw_text)
        flat = _Flattener(report, item_name, page, pending)
        title_heading: list = []
        _walk_html(root, flat, title_heading)
        flat.flush()
        # Page numbering runs on across the spine documents.
        page, pending = flat.page, flat.pending_pages
        if not flat.blocks and not flat.notes:
            continue
        fallback = toc_titles.get(item_name) or toc_titles.get(Path(item_name).name) or f"Capítulo {index}"
        chapter = _chapter_from_blocks(flat, fallback, "epub")
        if not chapter.paragraphs and chapter.title_generated:
            continue
        for source in chapter.sources:
            source["document"] = item_name
        result.chapters.append(chapter)
    if not result.chapters:
        result.chapters = [BookChapter(result.title, _paragraphs(fallback_text), title_generated=True)]


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def _from_docx(content: bytes, result: BookStructure, fallback_text: str, report: ExtractionReport,
               raw: list[str]) -> None:
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(content))
    report.method = "docx"
    result.title = document.core_properties.title or result.title
    result.author = document.core_properties.author or ""
    raw.append(" ".join(node.text or "" for node in document.element.body.iter(qn("w:t"))))
    # Word keeps footnotes and endnotes in parts of their own. They were never
    # read, and the word count above did not see them either, so a note's text
    # disappeared without a warning. They are read at the end of the chapter
    # that refers to them, "Nota 1. …", like the notes of every other format.
    note_texts = {kind: _docx_notes(document, kind) for kind in ("footnote", "endnote")}
    for texts in note_texts.values():
        raw.extend(texts.values())
    note_numbers: dict[tuple[str, str], int] = {}
    chapters: list[BookChapter] = []
    active = BookChapter(result.title, title_generated=True)
    active_notes: list[str] = []
    pending_pages: list[str] = []

    def close_chapter() -> None:
        nonlocal active_notes
        for note in active_notes:
            active.paragraphs.append(note)
            active.sources.append({"element": "note", "method": "docx"})
        active_notes = []

    for element in document.element.body.iterchildren():
        if element.tag == qn("w:tbl"):
            # Tables used to be skipped entirely: python-docx's paragraph list
            # does not include them.
            for row in Table(element, document._body).rows:
                cells = []
                for cell in row.cells:
                    text = _spoken(cell.text)
                    if text and (not cells or cells[-1] != text):  # merged cells repeat
                        cells.append(text)
                if cells:
                    active.paragraphs.append("; ".join(cells))
                    source = {"element": "table-row", "method": "docx"}
                    if pending_pages:
                        source["page_breaks"], pending_pages = pending_pages, []
                    active.sources.append(source)
            continue
        if element.tag != qn("w:p"):
            continue
        paragraph = Paragraph(element, document._body)
        text, pages = _docx_paragraph_text(paragraph)
        for page in pages:
            report.drop("page_number", page, "Número de página del original (estilo Page Number DAISY): marca "
                        "la posición, no se narra.", f"página {page}")
        pending_pages.extend(pages)
        text = _spoken(text)
        for kind, reference in _docx_note_references(element):
            body = note_texts[kind].get(reference, "")
            if body and (kind, reference) not in note_numbers:
                note_numbers[(kind, reference)] = len(note_numbers) + 1
                active_notes.append(note_text(f"{note_numbers[(kind, reference)]}. {body}"))
        if not text:
            continue
        level = _docx_heading_level(paragraph)
        page_source = {"page_breaks": pending_pages} if pending_pages else {}
        if level == 1:
            close_chapter()
            if active.paragraphs:
                chapters.append(active)
            active = BookChapter(text, title_source={"element": "heading", "heading_level": 1,
                                                     "method": "docx", **page_source})
        elif level:
            # Título 2 and below stay in place and are read where they are; the
            # level lets the audiobook's outline nest them under the chapter.
            active.paragraphs.append(text)
            active.sources.append({"element": "heading", "heading_level": level, "method": "docx", **page_source})
        else:
            for index, part in enumerate(_paragraphs(text)):
                active.paragraphs.append(part)
                active.sources.append({"element": "p", "method": "docx", **(page_source if index == 0 else {})})
        pending_pages = []
    close_chapter()
    if active.paragraphs or not chapters:
        chapters.append(active)
    result.chapters = chapters
    for section in document.sections:
        for part, label in ((section.header, "cabecera"), (section.footer, "pie de página")):
            try:
                text = " ".join(p.text for p in part.paragraphs)
            except Exception:
                continue
            report.drop("page_margin", text, f"Texto de {label} del documento; no se narra.", label)
    if not any(chapter.paragraphs for chapter in result.chapters):
        result.chapters = [BookChapter(result.title, _paragraphs(fallback_text), title_generated=True)]


_DOCX_RELATIONSHIPS = {
    "footnote": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
    "endnote": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
}
_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_notes(document, kind: str) -> dict[str, str]:
    """Footnote or endnote id -> its text. Word's separator entries are left out."""
    for relationship in document.part.rels.values():
        if relationship.reltype != _DOCX_RELATIONSHIPS[kind] or relationship.is_external:
            continue
        try:
            root = ET.fromstring(relationship.target_part.blob)
        except ET.ParseError:
            return {}
        notes: dict[str, str] = {}
        for note in root.iter(f"{_WORD_NS}{kind}"):
            if note.get(f"{_WORD_NS}type") in {"separator", "continuationSeparator", "continuationNotice"}:
                continue
            paragraphs = ["".join(node.text or "" for node in paragraph.iter(f"{_WORD_NS}t"))
                          for paragraph in note.iter(f"{_WORD_NS}p")]
            text = _spoken(" ".join(paragraph for paragraph in paragraphs if paragraph.strip()))
            if text:
                notes[note.get(f"{_WORD_NS}id", "")] = text
        return notes
    return {}


def _docx_note_references(paragraph_element) -> list[tuple[str, str]]:
    """The footnotes and endnotes a paragraph refers to, in order."""
    references = []
    for node in paragraph_element.iter(f"{_WORD_NS}footnoteReference", f"{_WORD_NS}endnoteReference"):
        kind = "footnote" if node.tag.endswith("footnoteReference") else "endnote"
        references.append((kind, node.get(f"{_WORD_NS}id", "")))
    return references


def _is_daisy_page_number(style) -> bool:
    """Save as DAISY's "Page Number DAISY" style marks a page number of the original."""
    if style is None:
        return False
    names = {getattr(style, "name", "") or "", getattr(style, "style_id", "") or ""}
    return any(re.sub(r"[^a-z]", "", name.casefold()) == "pagenumberdaisy" for name in names)


def _docx_paragraph_text(paragraph) -> tuple[str, list[str]]:
    """A paragraph's text without its page numbers, and those page numbers.

    Only a paragraph that uses the Save as DAISY page-number style is taken
    apart run by run; every other paragraph is read exactly as before.
    """
    if _is_daisy_page_number(paragraph.style):
        return "", [label for label in [_spoken(paragraph.text)] if label]
    try:
        items = list(paragraph.iter_inner_content())
    except AttributeError:  # an older python-docx
        return paragraph.text, []
    if not any(_is_daisy_page_number(getattr(item, "style", None)) for item in items if hasattr(item, "style")):
        return paragraph.text, []
    pages: list[str] = []
    parts: list[str] = []
    for item in items:
        if hasattr(item, "style") and _is_daisy_page_number(item.style):
            label = _spoken(item.text)
            if label:
                pages.append(label)
            continue
        parts.append(item.text)
    return "".join(parts), pages


def _docx_heading_level(paragraph) -> int | None:
    """1 for Título 1 / Heading 1, 2 for Título 2, and so on; None for body text."""
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = re.fullmatch(r"(?:heading|t[ií]tulo)\s+([1-9])", name.strip().lower())
    return int(match.group(1)) if match else None


# --------------------------------------------------------------------------- #
# PDF: text layer by layout, OCR for pages without one
# --------------------------------------------------------------------------- #

# Fewer characters than this and a page has no usable text layer.
_MIN_PAGE_CHARS = 25


def _page_profile(page, number: int) -> dict:
    text = page.get_text("text")
    area = abs(page.rect) or 1.0
    image_area = 0.0
    try:
        for info in page.get_image_info():
            x0, y0, x1, y1 = info.get("bbox", (0, 0, 0, 0))
            image_area += max(0.0, (x1 - x0)) * max(0.0, (y1 - y0))
    except Exception:
        image_area = 0.0
    chars = len(text.strip())
    ratio = min(1.0, image_area / area)
    if chars >= _MIN_PAGE_CHARS:
        status = "mixed" if ratio >= 0.5 and chars < 400 else "text"
    elif ratio > 0.05:
        status = "image_only"
    else:
        status = "blank"
    return {"page": number, "chars": chars, "image_ratio": round(ratio, 3), "status": status, "_text": text}


def _from_pdf(content: bytes, result: BookStructure, fallback_text: str, pages: str | None,
              report: ExtractionReport, raw: list[str]) -> None:
    with fitz.open(stream=content, filetype="pdf") as pdf:
        span = parse_page_range(pages, pdf.page_count)
        metadata = pdf.metadata or {}
        result.title = (metadata.get("title") or "").strip() or result.title
        result.author = (metadata.get("author") or "").strip()
        first, last = span or (1, pdf.page_count)
        profiles = [_page_profile(pdf[number - 1], number) for number in range(first, last + 1)]
        ocr_ready = ocr.available("spa")
        ocr_pages: dict[int, ocr.OcrPage] = {}
        for profile in profiles:
            number, status = profile["page"], profile["status"]
            if status == "image_only":
                if ocr_ready:
                    page_result = ocr.ocr_pdf_page(pdf[number - 1], number)
                    if page_result.error or not page_result.paragraphs:
                        report.issue(BLOCKING, "ocr_failed", f"La página {number} es una imagen y el OCR no "
                                     f"pudo leerla: {page_result.error or 'sin texto reconocido'}.",
                                     location=f"página {number}")
                    else:
                        ocr_pages[number] = page_result
                        profile["method"] = "ocr"
                        profile["ocr"] = page_result.as_dict()
                        raw.append("\n\n".join(page_result.paragraphs))
                        if page_result.mean_confidence < 80 or page_result.low_confidence_words:
                            report.issue(WARNING, "ocr_low_confidence",
                                         f"OCR de la página {number}: confianza media "
                                         f"{page_result.mean_confidence:.0f}/100 y "
                                         f"{len(page_result.low_confidence_words)} palabras dudosas.",
                                         location=f"página {number}",
                                         text=", ".join(item["word"] for item in page_result.low_confidence_words[:15]))
                else:
                    report.issue(BLOCKING, "ocr_required",
                                 f"La página {number} es una imagen sin capa de texto y no hay OCR "
                                 "disponible (instala tesseract-ocr con el modelo spa). Su contenido "
                                 "falta en el audiolibro.", location=f"página {number}")
            elif status == "mixed":
                profile["method"] = "text"
                raw.append(profile["_text"])
                if ocr_ready:
                    page_result = ocr.ocr_pdf_page(pdf[number - 1], number)
                    embedded = set(extraction_qa.words(profile["_text"]))
                    extra = [word for word in extraction_qa.words("\n".join(page_result.paragraphs))
                             if word not in embedded]
                    profile["ocr_extra_words"] = len(extra)
                    if len(extra) >= 15:
                        report.issue(WARNING, "text_in_image",
                                     f"La página {number} contiene una imagen con unas {len(extra)} palabras "
                                     "que no están en la capa de texto. Revisa si forman parte del libro.",
                                     location=f"página {number}", text=" ".join(extra[:25]))
                else:
                    report.issue(WARNING, "image_unchecked",
                                 f"La página {number} es sobre todo una imagen con poco texto; sin OCR no se "
                                 "puede comprobar si la imagen contiene texto.", location=f"página {number}")
            elif status == "text":
                profile["method"] = "text"
                raw.append(profile["_text"])
            else:
                profile["method"] = "none"
        report.pages = [{key: value for key, value in profile.items() if not key.startswith("_")}
                        for profile in profiles]
        result.chapters = _pdf_layout_chapters(pdf, span, ocr_pages=ocr_pages, report=report)
        if result.chapters:
            report.method = "pdf-layout" + ("+ocr" if ocr_pages else "")
        else:
            report.method = "pdf-pages" + ("+ocr" if ocr_pages else "")
            result.chapters = _pdf_page_chapters(profiles, ocr_pages, report)
    if not result.chapters:
        result.chapters = [BookChapter(result.title, _paragraphs(fallback_text), title_generated=True)]


def _pdf_page_chapters(profiles: list[dict], ocr_pages: dict[int, ocr.OcrPage],
                       report: ExtractionReport) -> list[BookChapter]:
    """One section per page, for PDFs without larger headings to go by.

    Running heads and page numbers repeated at the page edges are removed (and
    recorded); a paragraph cut by the page break is put back together on the
    page where it starts.
    """
    page_lines = {profile["page"]: [line.strip() for line in profile["_text"].splitlines() if line.strip()]
                  for profile in profiles if profile["page"] not in ocr_pages}
    repeated = extraction_qa.detect_repeated_margin_lines(list(page_lines.values()))
    chapters: list[BookChapter] = []
    for profile in profiles:
        number = profile["page"]
        if number in ocr_pages:
            paragraphs = list(ocr_pages[number].paragraphs)
            method = "ocr"
        elif number in page_lines and page_lines[number]:
            lines = page_lines[number]
            kept = []
            for position, line in enumerate(lines):
                at_edge = position == 0 or position == len(lines) - 1
                if at_edge and len(lines) > 1 and re.fullmatch(r"[\divxlc\s\-–—.]+", line, re.I) and len(line) <= 8:
                    report.drop("page_number", line, "Número de página del original; no se narra.", f"página {number}")
                    continue
                if at_edge and len(lines) > 1 and extraction_qa.is_repeated_margin(line, repeated):
                    report.drop("running_head", line, "Cabecera o pie repetido en muchas páginas.", f"página {number}")
                    continue
                kept.append(line)
            paragraphs = _pdf_paragraphs("\n".join(kept)) if kept else []
            method = "text"
        else:
            continue
        if not paragraphs:
            continue
        previous = chapters[-1] if chapters else None
        if previous and previous.paragraphs and not _SENTENCE_END.search(previous.paragraphs[-1]) \
                and paragraphs[0][:1].islower():
            tail = previous.paragraphs[-1]
            previous.paragraphs[-1] = (tail[:-1] + paragraphs[0] if tail.endswith("-") and tail[-2:-1].isalpha()
                                       else f"{tail} {paragraphs[0]}")
            paragraphs = paragraphs[1:]
            if not paragraphs:
                continue
        chapters.append(BookChapter(f"Página {number}", paragraphs, title_generated=True,
                                    sources=[{"page": str(number), "method": method} for _ in paragraphs]))
    for chapter in chapters:
        if len(chapter.sources) != len(chapter.paragraphs):
            chapter.sources = [dict(chapter.sources[0]) if chapter.sources else {} for _ in chapter.paragraphs]
    mark_page_starts(chapters)
    return chapters


def _pdf_paragraphs(text: str) -> list[str]:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    has_index_heading = any(re.fullmatch(r"(?:índice(?: general)?|tabla de contenidos?|contenido|sumario|contents|index)", line, re.IGNORECASE) for line in lines)
    has_leaders = sum(bool(re.search(r"(?:\.{2,}|…{2,})\s*(?:[ivxlcdm]+|\d{1,4})\s*$", line, re.IGNORECASE)) for line in lines) >= 2
    if lines and (has_index_heading or has_leaders):
        return lines
    return _paragraphs(text)


def _epub_metadata(book, name: str) -> str:
    values = book.get_metadata("DC", name) or []
    return " ".join(str(values[0][0]).split()) if values else ""


def _paragraphs(text: str) -> list[str]:
    """Split into paragraphs, rejoining words a line break split in two.

    Books hyphenate a word that does not fit on a line. Collapsing the
    whitespace on its own left "distribui- dores", which a voice reads as two
    broken pieces and a listener hears as a stutter. Every format goes through
    here, so EPUB, DOCX, XML and plain text get the same repair the
    layout-aware PDF reader already did.

    The soft hyphen (U+00AD) is an invisible "you may break here" mark and is
    simply removed.
    """
    blocks = re.split(r"\n\s*\n", text.replace("­", ""))
    return [joined for block in blocks if block.strip()
            if (joined := _join_lines(block.splitlines()))]


_EXPLICIT_CHAPTER_HEADING = re.compile(
    r"^(?:"
    r"(?:cap[ií]tulo|chapter|chapitre|capitolo|kapitel|parte|part|secci[oó]n|section)"
    r"\s+(?:[ivxlcdm]+|\d+|primero|primera|segundo|segunda|tercero|tercera|"
    r"cuarto|cuarta|quinto|quinta|sexto|sexta|séptimo|séptima|octavo|octava|"
    r"noveno|novena|décimo|décima|one|two|three|four|five|six|seven|eight|nine|ten)"
    r"(?:\s*[:\-–—]\s*.*)?\.?"
    r")$",
    re.IGNORECASE,
)


# Placeholders a transcriber leaves where the printed book had a picture:
# "[Ilustración]", "[Illustration: …]". They describe the edition, they are not
# the author's text, and read aloud they interrupt the story.
_TRANSCRIBER_MARKER = re.compile(r"^\s*\[(?:ilustraci[oó]n|illustration|imagen|image|grabado|figura)\b[^\]]*\]\s*$",
                                 re.IGNORECASE | re.MULTILINE)


# Italics written the plain-text way, "_¡Relámpago!_": the underscores are
# typography, not words.
_UNDERSCORE_ITALIC = re.compile(r"(?<![\w_])_(?=\S)([^_\n]{1,200}?)(?<=\S)_(?![\w_])")


def _drop_transcriber_markers(text: str, report: ExtractionReport) -> str:
    for match in _TRANSCRIBER_MARKER.finditer(text):
        report.drop("transcriber_marker", match.group(0), "Marca de ilustración de la edición digital; no es texto "
                    "del autor y no se narra.", "texto")
    text = _TRANSCRIBER_MARKER.sub("", text)
    italics = len(_UNDERSCORE_ITALIC.findall(text))
    if italics:
        report.issue(INFO, "italic_markup", f"Se quitaron los guiones bajos de {italics} cursivas escritas como "
                     "_texto_; las palabras se conservan.")
    return _UNDERSCORE_ITALIC.sub(r"\1", text)


# A section number standing alone ("I", "II.", "XIV") in a plain-text book.
_ROMAN_SECTION = re.compile(r"^[IVXLC]{1,7}\.?$")


def _is_section_heading(line: str) -> bool:
    from app.spanish_speech import roman_to_int

    return bool(_EXPLICIT_CHAPTER_HEADING.fullmatch(line)) or (
        bool(_ROMAN_SECTION.fullmatch(line)) and roman_to_int(line.rstrip(".")) is not None)


def _split_text_chapters(book_title: str, text: str) -> list[BookChapter]:
    """Split plain text only at explicit, standalone chapter/part headings."""
    blocks: list[tuple[str, bool]] = []
    # The soft hyphen only marks where a break is allowed; it is never spoken.
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "")
    for block in re.split(r"\n\s*\n", text):
        lines = [re.sub(r"\s+", " ", line).strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        # Preserve line boundaries in contents lists so the narration planner
        # can recognize leaders/page references instead of seeing one merged blob.
        structural = len(lines) > 1 and any(
            _EXPLICIT_CHAPTER_HEADING.fullmatch(line)
            or re.fullmatch(r"(?:indice(?: general)?|tabla de contenidos?|contenido|sumario|contents|index)", line, re.IGNORECASE)
            or re.search(r"(?:\.{2,}|…{2,})\s*(?:[ivxlcdm]+|\d{1,4})\s*$", line, re.IGNORECASE)
            for line in lines
        )
        if structural:
            blocks.extend((line, _is_section_heading(line)) for line in lines)
            continue
        # A heading must be an entire paragraph, never merely words at the
        # start of a prose paragraph. That avoids inventing chapter breaks.
        # _join_lines rather than " ".join: a word the line break hyphenated in
        # two ("distribui-" + "dores") is one word, and a voice reads the two
        # pieces separately.
        paragraph = _join_lines(lines)
        is_heading = len(lines) == 1 and _is_section_heading(lines[0])
        blocks.append((paragraph, is_heading))

    heading_indexes = [index for index, (_, is_heading) in enumerate(blocks) if is_heading]
    if len(heading_indexes) < 2:
        return [BookChapter(book_title, [paragraph for paragraph, _ in blocks], title_generated=True)]

    chapters: list[BookChapter] = []
    preface = [paragraph for paragraph, is_heading in blocks[:heading_indexes[0]] if not is_heading]
    if preface:
        # The book does not say "Prólogo": the label only lets the narration
        # plan classify what comes before chapter one. It is never narrated.
        chapters.append(BookChapter("Prólogo", preface, title_generated=True))
    active_title = ""
    active_paragraphs: list[str] = []
    for paragraph, is_heading in blocks[heading_indexes[0]:]:
        if is_heading:
            if active_title:
                chapters.append(BookChapter(active_title, active_paragraphs))
            active_title = paragraph
            active_paragraphs = []
        else:
            active_paragraphs.append(paragraph)
    if active_title:
        chapters.append(BookChapter(active_title, active_paragraphs))

    # Don't claim headings are chapters when the file only contains heading
    # labels with no readable chapter content.
    if len(chapters) < 2 or not any(chapter.paragraphs for chapter in chapters):
        return [BookChapter(book_title, [paragraph for paragraph, _ in blocks], title_generated=True)]
    return chapters


def _spoken(text: str) -> str:
    """Normalise a block of source text into one line a voice can read.

    Collapses whitespace and rejoins a word the line break hyphenated in two.
    Every format funnels through here, so XML, EPUB, DOCX and plain text all
    reach the voice with whole words.
    """
    return _join_lines(text.replace("­", "").splitlines())


def _text(node: ET.Element | None) -> str:
    return _spoken("".join(node.itertext())) if node is not None else ""


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _ancestors(node: ET.Element, parents: dict[ET.Element, ET.Element]):
    current = parents.get(node)
    while current is not None:
        yield current
        current = parents.get(current)


# --------------------------------------------------------------------------- #
# PDF: read the page layout, not just its text
# --------------------------------------------------------------------------- #
#
# get_text("text") gives one undivided run of lines per page. Narrated as is,
# that reads the running header and the page number aloud on every page,
# splits words hyphenated at line ends ("distribui- dores", which a voice says
# as two broken pieces and sounds like a stutter), and cuts every sentence that
# continues on the next page. The layout says which line is which: running
# heads and page numbers sit in the margins, chapter titles are set larger than
# the body, footnotes smaller and at the foot of the page.

_SENTENCE_END = re.compile(r"[.!?…:;»”\")\]]\s*$")
_FOOTNOTE_START = re.compile(r"^\s*(\d{1,3})[\s\t.)]+(.*)$", re.S)


def parse_page_range(value: str | None, page_count: int) -> tuple[int, int] | None:
    """``"20-114"`` -> (20, 114), clamped to the document; blank -> None."""
    text = (value or "").strip()
    if not text:
        return None
    match = re.fullmatch(r"(\d+)\s*(?:[-–—a]|hasta)?\s*(\d+)?", text, flags=re.I)
    if not match:
        raise ValueError("Escribe las páginas como «20-114».")
    first = int(match.group(1))
    last = int(match.group(2) or match.group(1))
    if first < 1 or last < first:
        raise ValueError("El rango de páginas no es válido: la primera página debe ser menor que la última.")
    if first > page_count:
        raise ValueError(f"El documento solo tiene {page_count} páginas.")
    return first, min(last, page_count)


def _line_text(spans: list[dict], footnote_numbers: set[str], dropped: list[str] | None = None) -> str:
    """A line's text without the footnote markers glued to the word before them."""
    parts: list[str] = []
    for index, span in enumerate(spans):
        text = span["text"]
        marker = text.strip()
        previous = parts[-1] if parts else ""
        if (marker in footnote_numbers and index > 0 and previous and not previous[-1].isspace()
                and not previous[-1].isdigit()):
            if dropped is not None:
                dropped.append(marker)
            continue
        parts.append(text)
    return "".join(parts)


def _join_lines(lines: list[str]) -> str:
    """Join a block's lines, undoing the hyphenation of words split at a line end."""
    text = ""
    for line in (value.strip() for value in lines):
        if not line:
            continue
        if text.endswith("-") and len(text) > 1 and text[-2].isalpha():
            # "discapa-" + "cidad" is one word; "Covid-" + "19" keeps its hyphen.
            text = text[:-1] + line if line[:1].islower() else text + line
        else:
            text = f"{text} {line}" if text else line
    return re.sub(r"\s+", " ", text).strip()


def _heading_text(lines: list[str]) -> str:
    """"Capítulo 1" over "El arte de hablar..." reads as "Capítulo 1. El arte de hablar..."."""
    parts = [line.strip() for line in lines if line.strip()]
    if len(parts) > 1 and _EXPLICIT_CHAPTER_HEADING.fullmatch(parts[0]):
        return f"{parts[0]}. {_join_lines(parts[1:])}"
    return _join_lines(parts)


def _pdf_layout_chapters(pdf, pages: tuple[int, int] | None, *, ocr_pages: dict | None = None,
                         report: ExtractionReport | None = None) -> list[BookChapter]:
    first, last = pages or (1, pdf.page_count)
    ocr_pages = ocr_pages or {}
    blocks = []  # (page, kind-independent data)
    sizes: dict[float, int] = {}
    for number in range(first, last + 1):
        if number in ocr_pages:
            # A scanned page has no font sizes; its paragraphs are body text.
            for position, paragraph in enumerate(ocr_pages[number].paragraphs):
                blocks.append({"page": number, "top": 0.1 + position * 0.001, "bottom": 0.1, "size": None,
                               "lines": [[{"text": paragraph, "size": None}]], "ocr": True})
            continue
        page = pdf[number - 1]
        height = page.rect.height
        for block in page.get_text("dict", sort=True)["blocks"]:
            if block.get("type") != 0 or not block.get("lines"):
                continue
            lines = [line["spans"] for line in block["lines"] if line["spans"]]
            text_size = [(span["size"], len(span["text"].strip())) for spans in lines for span in spans]
            chars = sum(count for _size, count in text_size)
            if not chars:
                continue
            size = max(text_size, key=lambda item: item[1])[0]
            for value, count in text_size:
                sizes[round(value * 2) / 2] = sizes.get(round(value * 2) / 2, 0) + count
            blocks.append({"page": number, "top": block["bbox"][1] / height, "bottom": block["bbox"][3] / height,
                           "size": size, "lines": lines})
    if not blocks or not sizes:
        return []
    body = max(sizes, key=sizes.get)
    for block in blocks:
        if block.get("ocr"):
            block["size"] = body
    if not any(block["size"] >= body + 2.5 for block in blocks):
        return []  # no headings to go by: keep the page-by-page reading

    # Footnote numbers per page, so their markers can be taken out of the text.
    for block in blocks:
        raw = "".join(span["text"] for span in block["lines"][0])
        match = _FOOTNOTE_START.match(raw)
        block["kind"] = "body"
        if block.get("ocr"):
            continue
        if block["top"] < 0.09 and block["size"] < body:
            block["kind"] = "margin"  # running head
            block["reason"] = ("running_head", "Cabecera de página (titulillo); no se narra.")
        elif block["top"] > 0.9 and re.fullmatch(r"[\divxlc\s]+", "".join(
                span["text"] for spans in block["lines"] for span in spans), re.I):
            block["kind"] = "margin"  # page number
            block["reason"] = ("page_number", "Número de página del original; no se narra.")
        elif block["size"] <= body - 1.25 and block["top"] > 0.6 and match:
            block["kind"] = "footnote"
            block["number"] = match.group(1)
        elif block["size"] >= body + 2.5:
            block["kind"] = "heading"
    notes_on_page: dict[int, set[str]] = {}
    for block in blocks:
        if block["kind"] == "footnote":
            notes_on_page.setdefault(block["page"], set()).add(block["number"])

    chapters: list[BookChapter] = []
    notes: list[tuple[str, int]] = []
    current = BookChapter("", title_generated=True)
    pending_heading: list[str] = []

    def close_chapter() -> None:
        nonlocal current, notes
        if notes:
            for note, page_number in notes:
                current.paragraphs.append(note)
                current.sources.append({"page": str(page_number), "element": "footnote", "method": "pdf-layout"})
            notes = []
        if current.paragraphs or current.title:
            chapters.append(current)

    for index, block in enumerate(blocks):
        numbers = notes_on_page.get(block["page"], set())
        markers: list[str] = []
        lines = [_line_text(spans, numbers, markers) for spans in block["lines"]]
        if report is not None:
            for marker in markers:
                report.drop("note_marker", marker, "Marca de llamada a nota; la nota se lee al final del capítulo.",
                            f"página {block['page']}")
        if block["kind"] == "margin":
            if report is not None:
                kind, reason = block["reason"]
                report.drop(kind, " ".join(lines), reason, f"página {block['page']}")
            continue
        if block["kind"] == "heading":
            pending_heading.extend(lines)
            following = blocks[index + 1] if index + 1 < len(blocks) else None
            if following and following["kind"] == "heading" and following["page"] == block["page"]:
                continue  # a title set over several blocks
            close_chapter()
            current = BookChapter(_heading_text(pending_heading),
                                  title_source={"element": "heading", "heading_level": 1,
                                                "page": str(block["page"]), "method": "pdf-layout"})
            pending_heading = []
            continue
        if block["kind"] == "footnote":
            text = _join_lines(lines)
            match = _FOOTNOTE_START.match(text)
            if match:
                notes.append((f"Nota {match.group(1)}. {match.group(2).strip()}", block["page"]))
                if report is not None:
                    report.drop("note_marker", match.group(1), "Número de la nota, leído como «Nota N».",
                                f"página {block['page']}")
            continue
        text = _join_lines(lines)
        if not text:
            continue
        previous = current.paragraphs[-1] if current.paragraphs else ""
        continues = (previous and block is _first_body_on_page(blocks, block["page"])
                     and not _SENTENCE_END.search(previous))
        if continues:
            # The paragraph carries on from the previous page.
            if previous.endswith("-") and text[:1].islower():
                current.paragraphs[-1] = previous[:-1] + text
            else:
                current.paragraphs[-1] = f"{previous} {text}"
        else:
            current.paragraphs.append(text)
            current.sources.append({"page": str(block["page"]),
                                    "method": "ocr" if block.get("ocr") else "pdf-layout"})
    close_chapter()
    result = [chapter for chapter in chapters if chapter.paragraphs]
    for chapter in result:
        if not chapter.title:
            chapter.title_generated = True
    mark_page_starts(result)
    return result


def mark_page_starts(chapters: list[BookChapter]) -> None:
    """Note where each page begins, for sources that only say which page they are on.

    A PDF paragraph knows its page, so a page begins at the first paragraph
    (or heading) found on it. A page that only continues a paragraph from the
    one before has no paragraph of its own to begin at and gets no marker:
    placing one elsewhere would point at the wrong text.
    """
    previous = ""
    for chapter in chapters:
        entries = ([chapter.title_source] if chapter.title_source and not chapter.title_generated else []) + [
            source for source in chapter.sources]
        for source in entries:
            page = str(source.get("page") or "")
            if page and page != previous and source.get("element") not in {"footnote", "note"}:
                source["page_breaks"] = [page]
                previous = page


def _first_body_on_page(blocks: list[dict], page: int) -> dict | None:
    return next((block for block in blocks if block["page"] == page and block["kind"] == "body"), None)
