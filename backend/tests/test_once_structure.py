"""The ONCE structure: planilla, initial files, levels, pages, notes, closing phrase.

``MODELO ONCE.docx`` is the supervisor's model. The demo book in the corpus
(``el-faro-de-la-isla-almendra.xml``) has the model's chapters and pages, so a
package built from it can be compared with the model file by file. Offline:
the mock provider stands in for the voice; nothing here judges audio quality.
"""

from __future__ import annotations

import asyncio
import io
import json
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import RGBColor
from fastapi.testclient import TestClient

from app import daisy_dtd, main, once_structure, once_validation, planilla
from app import jobs as job_store
from app.audio_cache import AudioCache
from app.book_structure import extract_book_structure
from app.daisy_export import build_daisy_package, make_paragraphs
from app.daisy_validation import validate_daisy_package
from app.narration_plan import apply_narration_selection, build_narration_plan

CORPUS = Path(__file__).parent / "corpus"
DEMO = CORPUS / "el-faro-de-la-isla-almendra.xml"
DTBOOK = "{http://www.daisy.org/z3986/2005/dtbook/}"
SMIL = "{http://www.w3.org/2001/SMIL20/}"
NCX = "{http://www.daisy.org/z3986/2005/ncx/}"
# The demo's planilla, as MODELO ONCE.docx gives it.
MODEL = {
    "file2_page": "2", "index_page": "3", "index_depth": 1,
    "edition_line": "PRIMERA EDICIÓN MADRID, FEBRERO DE 2026", "publisher_line": "EDITORIAL CRÍTICA",
    "production_date": "octubre de 2026",
}
MP3 = (b"\xff\xfb\x90\x64" + b"\x00" * 413) * 40  # 1.04 s of real MP3 frames


def _structure(sheet: dict, source: bytes | None = None, name: str = "demo.xml"):
    structure = extract_book_structure(name, source if source is not None else DEMO.read_bytes())
    plan = build_narration_plan([(c.title, c.paragraphs, c.title_generated) for c in structure.chapters])
    selected = apply_narration_selection(plan, None)
    metadata = {"title": structure.title, "author": structure.author, "edition": structure.edition,
                "publication_year": structure.publication_year, "publisher": structure.publisher}
    return once_structure.build(selected, once_structure.selected_provenance(selected, structure), metadata,
                                planilla.from_dict(sheet))


def _package(built, tmp_path: Path) -> Path:
    chapters, layout = [], []
    for section, record in zip(built.sections, built.records()):
        texts = [paragraph.text for paragraph in section.paragraphs]
        body = ([section.title] if section.heading_spoken else []) + [part for text in texts for part in make_paragraphs(text)]
        chapters.append((section.title, body))
        layout.append({"kind": section.kind, "level": section.level,
                       "units": once_structure.unit_layout(record, texts, section.heading_spoken, make_paragraphs)})
    segments = [MP3 for _title, body in chapters for _text in body]
    return build_daisy_package(built.resolved.title, [text for _t, body in chapters for text in body], segments,
                               tmp_path / "package", chapters=chapters, author=built.resolved.author, layout=layout,
                               identifiers=[("ABSYS", built.resolved.planilla.absys)])


def _records(built) -> list[dict]:
    return [{"title": section.title, "paragraphs": [paragraph.text for paragraph in section.paragraphs],
             "title_generated": section.title_generated, "once": record}
            for section, record in zip(built.sections, built.records())]


# --------------------------------------------------------------------------- #
# The planilla
# --------------------------------------------------------------------------- #

def test_the_legal_notice_is_the_models_text_word_for_word() -> None:
    text = planilla.ONCE_LEGAL_TEXT
    assert text.startswith("Esta adaptación es propiedad de la Organización Nacional de Ciegos Españoles (ONCE) y ha "
                           "sido realizada en sistema Daisy exclusivamente para su reproducción por personas afiliadas")
    assert "conforme a lo dispuesto en la Ley española de Propiedad Intelectual." in text
    assert "al amparo de la normativa y los tratados internacionales aplicables en cada momento." in text
    # Kept as the model ends it, closing quotation mark included: not corrected silently.
    assert text.endswith("Servicio Bibliográfico ONCE 2026\u201d")
    assert len(text) == 961
    assert planilla.defaults()["legal_text"] == text


def test_the_models_invented_book_is_never_a_default() -> None:
    values = json.dumps(planilla.defaults(), ensure_ascii=False)
    for invented in ("El faro de la isla Almendra", "Antonio Pérez", "CRÍTICA", "FEBRERO DE 2026"):
        assert invented not in values
    resolved = planilla.resolve(planilla.Planilla(), {"title": "Otro libro", "author": "Otra autora"})
    assert (resolved.title, resolved.author) == ("Otro libro", "Otra autora")
    # What the book does not supply is reported, not invented.
    assert "editorial" in resolved.missing and "página del índice" in resolved.missing


def test_a_planilla_with_unknown_fields_or_bad_values_is_refused() -> None:
    with pytest.raises(ValueError):
        planilla.from_json('{"file2_page": "dos"}')
    with pytest.raises(ValueError):
        planilla.from_json('{"closing": "Fin de {libro}."}')
    with pytest.raises(ValueError):
        planilla.from_json('{"index_depth": 3}')
    assert planilla.from_json("") is None
    assert planilla.from_json('{"enabled": false}') is None


def _model_docx() -> bytes:
    """A planilla written the way MODELO ONCE.docx is: content in black, instructions in red."""
    document = Document()

    def line(*parts):
        paragraph = document.add_paragraph()
        for text, red in parts:
            run = paragraph.add_run(text)
            if red:
                run.font.color.rgb = RGBColor(0xEE, 0x00, 0x00)
                run.bold = True

    line(("Número ABSYS. ", False), ("Es el número que identifica el audiolibro. En esta demo no hay.", True))
    line(("ARCHIVO 1 DE AUDIO.", True))
    line(("El faro de la isla Almendra, ", False), ("(SIEMPRE SE PONE EL TÍTULO)", True))
    line(("De Antonio Pérez ", False), ("(TRAS UNA COMA EL AUTOR, LUEGO EL TEXTO NORMATIVO DE ONCE)", True))
    line((planilla.ONCE_LEGAL_TEXT, False))
    line(("EL ARCHIVO 1 ACABA SIEMPRE EN LA PÁGINA 1.", True))
    line(("ARCHIVO 2 DE AUDIO.", True))
    line(("DATOS BIBLIOGRÁFICOS", False))
    line(("El faro de la isla Almendra, ", False))
    line(("De Antonio Pérez ", False))
    line(("PRIMERA EDICIÓN MADRID,  FEBRERO DE 2026", False))
    line(("EDITORIAL CRÍTICA", False))
    line(("Obra grabada en formato digital para el Servicio Bibliográfico de la ONCE, Sonidos S-2 Productions, "
          "S.L. octubre de 2026", False))
    line(("EN ESTE CASO PONERLA EN LA PÁGINA 2", True))
    line(("ARCHIVO 3 DE AUDIO ", True))
    for text in ("Índice. ", "Índice pagina 3.", "Capítulo 1. La llegada. Pagina 4.", "Capítulo 2. La tormenta. Página 5.",
                 "Capitulo 3. El relevo. Página 7.", "Fin de índice."):
        line((text, False))
    for number, heading in ((4, "Capítulo 1. La llegada."), (5, "El inventario."), (6, "Capítulo 2.  La tormenta."),
                            (7, "La reparación."), (8, "Capitulo 3. El relevo.")):
        line((f"ARCHIVO {number} DE AUDIO.", True))
        line((heading, False))
    line(("AL FINAL DE CAPITULO 3 DEBE DECIR.", True))
    line(("Aquí finaliza la grabación de El faro de la isla Almendra, De Antonio Pérez.", False))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_planilla_written_like_the_model_is_read_field_by_field() -> None:
    values, warnings = planilla.parse_docx(_model_docx())
    assert values["title"] == "El faro de la isla Almendra"
    assert values["author"] == "Antonio Pérez"
    assert values["legal_text"] == planilla.ONCE_LEGAL_TEXT
    assert (values["file1_page"], values["file2_page"], values["index_page"]) == ("1", "2", "3")
    assert values["edition_line"] == "PRIMERA EDICIÓN MADRID, FEBRERO DE 2026"
    assert values["publisher_line"] == "EDITORIAL CRÍTICA"
    assert values["production_line"] == planilla.PRODUCTION_LINE
    assert values["production_date"] == "octubre de 2026"
    assert values["chapter_pages"] == {"Capítulo 1. La llegada": "4", "Capítulo 2. La tormenta": "5",
                                       "Capitulo 3. El relevo": "7"}
    assert [item["heading"] for item in values["expected_files"]] == [
        "Capítulo 1. La llegada", "El inventario", "Capítulo 2. La tormenta", "La reparación", "Capitulo 3. El relevo"]
    # The model's index lists only the chapters, not «El inventario» or «La reparación».
    assert values["index_depth"] == 1
    assert values["closing"] == "Aquí finaliza la grabación de {title_and_author}."
    assert values["absys"] == "" and any("ABSYS" in warning for warning in warnings)
    planilla.from_dict(values)  # and it is a valid planilla


@pytest.mark.skipif(not Path("D:/Downloads/MODELO ONCE.docx").is_file(), reason="the supervisor's file is not here")
def test_the_supervisors_own_model_file_is_read() -> None:
    values, _warnings = planilla.parse_docx(Path("D:/Downloads/MODELO ONCE.docx").read_bytes())
    assert values["legal_text"] == planilla.ONCE_LEGAL_TEXT
    assert values["file2_page"] == "2" and values["index_page"] == "3" and values["index_depth"] == 1


def test_the_parse_endpoint_reads_a_planilla_and_refuses_other_files() -> None:
    client = TestClient(main.app)
    response = client.post("/api/planilla/parse", files={"file": ("MODELO ONCE.docx", io.BytesIO(_model_docx()),
                                                                  "application/octet-stream")})
    assert response.status_code == 200, response.text
    assert response.json()["planilla"]["file2_page"] == "2"
    assert client.post("/api/planilla/parse", files={"file": ("x.txt", io.BytesIO(b"x"), "text/plain")}).status_code == 415
    defaults = client.get("/api/planilla/defaults").json()
    assert defaults["planilla"]["legal_text"] == planilla.ONCE_LEGAL_TEXT
    assert re.fullmatch(r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)"
                        r" de \d{4}", defaults["production_date"])


# --------------------------------------------------------------------------- #
# The structure, from the book
# --------------------------------------------------------------------------- #

def test_the_demo_book_has_the_models_eight_files() -> None:
    built = _structure(MODEL)
    assert [(section.kind, section.level, section.title) for section in built.sections] == [
        ("once_title", 1, "El faro de la isla Almendra, De Antonio Pérez"),
        ("once_bibliographic", 1, "DATOS BIBLIOGRÁFICOS"),
        ("once_index", 1, "Índice."),
        ("chapter", 1, "Capítulo 1. La llegada"),
        ("subchapter", 2, "El inventario"),
        ("chapter", 1, "Capítulo 2. La tormenta"),
        ("subchapter", 2, "La reparación"),
        ("chapter", 1, "Capítulo 3. El relevo"),
    ]
    first, second, index = built.sections[:3]
    assert [paragraph.text for paragraph in first.paragraphs] == [planilla.ONCE_LEGAL_TEXT]
    assert first.title_pages == ["1"] and second.title_pages == ["2"] and index.title_pages == ["3"]
    assert [paragraph.text for paragraph in second.paragraphs] == [
        "El faro de la isla Almendra, De Antonio Pérez", "PRIMERA EDICIÓN MADRID, FEBRERO DE 2026", "EDITORIAL CRÍTICA",
        "Obra grabada en formato digital para el Servicio Bibliográfico de la ONCE, Sonidos S-2 Productions, S.L. "
        "octubre de 2026"]
    # The model's index, generated from the book: its pages are the book's own.
    assert [paragraph.text for paragraph in index.paragraphs] == [
        "Índice página 3.", "Capítulo 1. La llegada. Página 4.", "Capítulo 2. La tormenta. Página 5.",
        "Capítulo 3. El relevo. Página 7.", "Fin de índice."]
    assert built.warnings == []


def test_an_index_of_depth_two_lists_the_subchapters_with_their_pages() -> None:
    built = _structure({**MODEL, "index_depth": 2})
    entries = [paragraph.text for paragraph in built.sections[2].paragraphs]
    assert entries == ["Índice página 3.", "Capítulo 1. La llegada. Página 4.", "El inventario. Página 4.",
                       "Capítulo 2. La tormenta. Página 5.", "La reparación. Página 6.",
                       "Capítulo 3. El relevo. Página 7.", "Fin de índice."]


def test_a_note_is_framed_once_and_the_closing_phrase_ends_the_last_chapter() -> None:
    built = _structure(MODEL)
    notes = [paragraph for section in built.sections for paragraph in section.paragraphs if paragraph.role == "note"]
    assert [note.text for note in notes] == [
        "Nota 1. En la isla, los faros todavía se encendían a mano cuando fallaba el generador. Fin de nota 1."]
    assert notes[0].note == 1
    every = [paragraph for section in built.sections for paragraph in section.paragraphs]
    closing = [paragraph for paragraph in every if paragraph.role == "closing"]
    assert [paragraph.text for paragraph in closing] == [
        "Aquí finaliza la grabación de El faro de la isla Almendra, De Antonio Pérez."]
    assert every[-1] is closing[0]


def test_a_closing_phrase_or_note_end_already_in_the_book_is_not_repeated() -> None:
    source = DEMO.read_text(encoding="utf-8").replace(
        "la luz no se apagó ninguna noche.</p>",
        "la luz no se apagó ninguna noche.</p><p>Aquí finaliza la grabación de El faro de la isla Almendra, "
        "De Antonio Pérez.</p>").replace(
        "fallaba el generador.</p>", "fallaba el generador. Fin de nota 1.</p>")
    built = _structure(MODEL, source.encode("utf-8"))
    texts = [paragraph.text for section in built.sections for paragraph in section.paragraphs]
    assert sum(text.startswith("Aquí finaliza la grabación") for text in texts) == 1
    assert [paragraph.role for paragraph in built.sections[-1].paragraphs][-1] == "closing"
    note = next(text for text in texts if text.startswith("Nota 1."))
    assert note.count("Fin de nota 1.") == 1


def test_a_book_without_page_numbers_takes_its_index_pages_from_the_planilla(tmp_path) -> None:
    text = ("Capítulo 1\n\nTexto del primer capítulo, largo y claro.\n\nCapítulo 2\n\n"
            "Texto del segundo capítulo, también claro.\n").encode("utf-8")

    def entries(sheet):
        built = _structure(sheet, text, "libro.txt")
        pages = [warning["code"] for warning in built.warnings if warning["code"] != "missing"]
        return built, [paragraph.text for paragraph in built.sections[2].paragraphs], pages

    # No pages anywhere: the index says none rather than inventing them (and
    # never gives a chapter the index's own page).
    built, listed, warnings = entries(MODEL)
    assert listed == ["Índice página 3.", "Capítulo 1.", "Capítulo 2.", "Fin de índice."]
    assert warnings == ["index_without_page"]
    # The planilla's index pages: each chapter starts there, with its marker.
    built, listed, warnings = entries({**MODEL, "chapter_pages": {"Capítulo 1": "4", "Capítulo 2": "6"}})
    assert listed == ["Índice página 3.", "Capítulo 1. Página 4.", "Capítulo 2. Página 6.", "Fin de índice."]
    assert [section.title_pages for section in built.sections[3:]] == [["4"], ["6"]]
    assert warnings == []
    # Only one page given: the next chapter is on it as far as the markers go,
    # and the outline says nobody stated it.
    built, listed, warnings = entries({**MODEL, "chapter_pages": {"Capítulo 1": "4"}})
    assert listed[2] == "Capítulo 2. Página 4."
    assert warnings == ["page_inherited"]
    # A plain-text book names no author, and that is reported, not invented.
    assert any(warning["code"] == "missing" and "autor" in warning["message"] for warning in built.warnings)
    # And the index agrees with the page markers the package really has.
    package = _package(built, tmp_path)
    checks = {check["id"]: check["status"] for check in
              once_validation.verify(package, _records(built), built.resolved.as_dict())}
    assert checks["once_index"] == "passed" and checks["once_pages"] == "passed"


def test_a_book_that_differs_from_the_planillas_files_is_reported() -> None:
    sheet = {**MODEL, "expected_files": [{"number": 4, "heading": "Capítulo 1. La llegada"},
                                         {"number": 5, "heading": "El faro apagado"}]}
    built = _structure(sheet)
    assert any(warning["code"] == "planilla_mismatch" and "El faro apagado" in warning["message"]
               for warning in built.warnings)


def test_a_subchapter_whose_chapter_has_nothing_to_read_becomes_a_chapter() -> None:
    source = DEMO.read_text(encoding="utf-8").replace("<h1>Capítulo 1. La llegada</h1>", "")
    source = re.sub(r'(<level1 id="cap1">.*?)<p>La barca.*?</p>\s*<p>El faro se alzaba.*?</p>', r"\1",
                    source, flags=re.S)
    built = _structure(MODEL, source.encode("utf-8"))
    inventory = next(section for section in built.sections if section.title == "El inventario")
    assert (inventory.kind, inventory.level) == ("chapter", 1)


# --------------------------------------------------------------------------- #
# The package and its checks
# --------------------------------------------------------------------------- #

def test_the_demo_package_passes_the_dtds_and_every_once_check(tmp_path) -> None:
    built = _structure(MODEL)
    package = _package(built, tmp_path)
    validate_daisy_package(package)
    assert daisy_dtd.validate_package(package)["status"] == "passed"
    checks = once_validation.verify(package, _records(built), built.resolved.as_dict())
    assert {check["id"]: check["status"] for check in checks} == {
        "once_initial_files": "passed", "once_hierarchy": "passed", "once_pages": "passed",
        "once_index": "passed", "once_notes": "passed", "once_closing": "passed", "once_play_order": "passed"}


def test_initial_files_are_front_matter_and_subchapters_nest(tmp_path) -> None:
    package = _package(_structure(MODEL), tmp_path)
    dtbook = ET.parse(next(package.glob("*.xml"))).getroot()
    front = dtbook.find(f"{DTBOOK}book/{DTBOOK}frontmatter")
    assert [level.get("class") for level in front.findall(f"{DTBOOK}level1")] == [
        "once-title", "once-bibliographic", "once-index"]
    body = dtbook.find(f"{DTBOOK}book/{DTBOOK}bodymatter")
    chapters = body.findall(f"{DTBOOK}level1")
    assert [level.find(f"{DTBOOK}h1").text for level in chapters] == [
        "Capítulo 1. La llegada", "Capítulo 2. La tormenta", "Capítulo 3. El relevo"]
    assert [level.find(f"{DTBOOK}h2").text for level in body.iter(f"{DTBOOK}level2")] == ["El inventario", "La reparación"]
    # One audio file per section: «ARCHIVO N» is audio/chapter-00N.
    assert sorted(path.parent.name for path in package.glob("audio/chapter-*/chapter-*.mp3")) == [
        f"chapter-{number:03d}" for number in range(1, 9)]
    opf = (next(package.glob("*.opf"))).read_text(encoding="utf-8")
    assert 'scheme="ABSYS"' not in opf  # the model has no ABSYS number


def test_page_numbers_are_silent_skippable_markers(tmp_path) -> None:
    package = _package(_structure(MODEL), tmp_path)
    dtbook = ET.parse(next(package.glob("*.xml"))).getroot()
    pages = list(dtbook.iter(f"{DTBOOK}pagenum"))
    assert [page.text for page in pages] == ["1", "2", "3", "4", "5", "6", "7"]
    for page in pages:
        smil_name, _, par_id = page.get("smilref").partition("#")
        smil = ET.parse(package / smil_name).getroot()
        par = next(node for node in smil.iter(f"{SMIL}par") if node.get("id") == par_id)
        assert par.get("customTest") == "pagenum"
        assert par.find(f"{SMIL}audio") is None  # nothing is read for a page number
        assert par.find(f"{SMIL}text").get("src").endswith("#" + page.get("id"))
        assert any(test.get("id") == "pagenum" for test in smil.iter(f"{SMIL}customTest"))
    ncx = ET.parse(next(package.glob("*.ncx"))).getroot()
    assert [target.get("value") for target in ncx.iter(f"{NCX}pageTarget")] == ["1", "2", "3", "4", "5", "6", "7"]
    metas = {meta.get("name"): meta.get("content") for meta in ncx.iter(f"{NCX}meta")}
    assert (metas["dtb:depth"], metas["dtb:totalPageCount"], metas["dtb:maxPageNumber"]) == ("2", "7", "7")
    tests = {test.get("id"): test.get("bookStruct") for test in ncx.iter(f"{NCX}smilCustomTest")}
    assert tests == {"pagenum": "PAGE_NUMBER", "note": "NOTE"}


def test_a_note_is_a_skippable_structure_listed_for_navigation(tmp_path) -> None:
    package = _package(_structure(MODEL), tmp_path)
    dtbook = ET.parse(next(package.glob("*.xml"))).getroot()
    note = next(dtbook.iter(f"{DTBOOK}note"))
    smil_name = note.find(f"{DTBOOK}p").get("smilref").partition("#")[0]
    smil = ET.parse(package / smil_name).getroot()
    sequence = next(node for node in smil.iter(f"{SMIL}seq") if node.get("customTest") == "note")
    assert sequence.find(f"{SMIL}par/{SMIL}audio") is not None  # a note is read
    ncx = ET.parse(next(package.glob("*.ncx"))).getroot()
    assert [target.find(f"{NCX}navLabel/{NCX}text").text for target in ncx.iter(f"{NCX}navTarget")] == ["Nota 1"]


def test_play_order_follows_the_reading_order_one_container_at_a_time(tmp_path) -> None:
    package = _package(_structure(MODEL), tmp_path)
    ncx = ET.parse(next(package.glob("*.ncx"))).getroot()
    orders = sorted(int(node.get("playOrder")) for node in ncx.iter()
                    if node.tag in {f"{NCX}navPoint", f"{NCX}pageTarget", f"{NCX}navTarget"})
    assert orders == list(range(1, len(orders) + 1))
    # Page 1 comes before the title it marks; the chapter after its page.
    first_page = next(ncx.iter(f"{NCX}pageTarget"))
    first_point = next(ncx.iter(f"{NCX}navPoint"))
    assert int(first_page.get("playOrder")) < int(first_point.get("playOrder"))


def _rewrite(path: Path, old: str, new: str) -> None:
    text = path.read_text(encoding="utf-8")
    assert old in text
    path.write_text(text.replace(old, new, 1), encoding="utf-8")


def test_the_checks_catch_a_package_that_is_not_what_the_planilla_says(tmp_path) -> None:
    built = _structure(MODEL)
    package = _package(built, tmp_path)
    dtbook = next(package.glob("*.xml"))
    tampered = tmp_path / "tampered"
    shutil.copytree(package, tampered)
    tampered_book = tampered / dtbook.name
    _rewrite(tampered_book, "Propiedad Intelectual", "Propiedad Industrial")
    _rewrite(tampered_book, "Capítulo 2. La tormenta. Página 5.", "Capítulo 2. La tormenta. Página 9.")
    _rewrite(tampered_book, 'class="once-closing"', 'class="once-text"')
    smil = tampered / "smil" / "chapter-0004.smil"
    _rewrite(smil, 'customTest="pagenum"><text src="../El_faro_de_la_isla_Almendra.xml#page-0004" />',
             'customTest="pagenum"><text src="../El_faro_de_la_isla_Almendra.xml#page-0004" />'
             '<audio src="../audio/chapter-004/chapter-004.mp3" clipBegin="0:00:00.000" clipEnd="0:00:00.500" />')
    statuses = {check["id"]: check for check in once_validation.verify(tampered, _records(built),
                                                                         built.resolved.as_dict())}
    assert statuses["once_initial_files"]["status"] == "failed"
    assert "texto legal" in statuses["once_initial_files"]["detail"]
    assert statuses["once_index"]["status"] == "failed"
    assert statuses["once_closing"]["status"] == "failed"
    assert statuses["once_pages"]["status"] == "failed"
    assert "silenciosa" in statuses["once_pages"]["detail"]


# --------------------------------------------------------------------------- #
# Through the application, offline
# --------------------------------------------------------------------------- #

@pytest.fixture
def offline(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "settings", main.settings.model_copy(update={
        "tts_provider": "mock", "asr_provider": "mock", "tts_voice_id": "mock-es-es-female",
        "asr_secondary_provider": "none", "quality_review_enabled": True,
        "once_intro_template": "", "once_bibliographic_template": "",
        "once_identification_statement": "", "once_legal_statement": "",
    }))
    monkeypatch.setattr(main, "storage_root", tmp_path / "data")
    monkeypatch.setattr(main, "database_path", tmp_path / "data" / "jobs.sqlite3")
    monkeypatch.setattr(main, "audio_cache", AudioCache(tmp_path / "data" / "audio-cache"))
    monkeypatch.setattr(main, "_start_or_resume_job", lambda job_id: asyncio.sleep(0))
    return tmp_path


def _demo_files():
    return {"file": ("el-faro-de-la-isla-almendra.xml", io.BytesIO(DEMO.read_bytes()), "application/xml")}


def test_a_reviewed_production_delivers_the_once_structure(offline) -> None:
    client = TestClient(main.app)
    response = client.post("/api/books/jobs", files=_demo_files(), data={"planilla_json": json.dumps(MODEL)})
    assert response.status_code == 202, response.text
    job_id = response.json()["job_id"]
    asyncio.run(main.run_production_job(job_id))
    job = job_store.get_job(main.database_path, job_id)
    assert job["state"] == "completed", job["error"]
    job_dir = Path(job["source_path"]).parent
    chain = json.loads((job_dir / "completeness.json").read_text(encoding="utf-8"))
    statuses = {check["id"]: check["status"] for check in chain["checks"]}
    # The mock voice tags each clip with its text, which is not MPEG audio, so
    # the audio-format check fails for simulated audio with or without a
    # planilla (and a mock book is never released). Every other link holds.
    assert {key: value for key, value in statuses.items() if key != "audio_format"} == {
        key: "passed" for key in statuses if key != "audio_format"}
    assert {key for key in statuses if key.startswith("once_")} == {
        "once_initial_files", "once_hierarchy", "once_pages", "once_index", "once_notes", "once_closing",
        "once_play_order"}
    # The legal notice and the closing phrase were narrated and checked like the rest.
    roles = [(unit.get("provenance") or {}).get("role") for unit in job["sentences"]]
    assert roles.count("closing") == 1 and "legal" in roles and "index_entry" in roles and "note" in roles
    assert all(unit["state"] == "verified" for unit in job["sentences"])
    with zipfile.ZipFile(job["output_path"]) as archive:
        names = archive.namelist()
        assert sum(1 for name in names if re.fullmatch(r"DAISY 3/audio/chapter-\d{3}/chapter-\d{3}\.mp3", name)) == 8
        book = archive.read(next(name for name in names if name.endswith(".xml") and name.count("/") == 1)).decode()
        assert book.count("Aquí finaliza la grabación de El faro de la isla Almendra, De Antonio Pérez.") == 1
        assert "Fin de nota 1." in book


def test_the_job_structure_gives_the_editor_files_units_pages_and_final_audio(offline) -> None:
    client = TestClient(main.app)
    job_id = client.post("/api/books/jobs", files=_demo_files(),
                         data={"planilla_json": json.dumps(MODEL)}).json()["job_id"]
    before = client.get(f"/api/jobs/{job_id}/structure").json()
    assert before["once"] is True and before["assembled"] is False
    assert [section["status"] for section in before["sections"]] == ["pending"] * 8
    asyncio.run(main.run_production_job(job_id))
    structure = client.get(f"/api/jobs/{job_id}/structure").json()
    assert [(section["file_number"], section["kind"], section["level"]) for section in structure["sections"]] == [
        (1, "once_title", 1), (2, "once_bibliographic", 1), (3, "once_index", 1), (4, "chapter", 1),
        (5, "subchapter", 2), (6, "chapter", 1), (7, "subchapter", 2), (8, "chapter", 1)]
    assert [section["page_start"] for section in structure["sections"]] == ["1", "2", "3", "4", "4", "5", "6", "7"]
    assert structure["pages"] == 7
    assert all(section["status"] == "verified" for section in structure["sections"])
    first = structure["sections"][0]
    assert [unit["role"] for unit in first["units"]] == ["heading", "legal", "legal"]
    assert first["units"][0]["page_breaks"] == ["1"]
    assert first["units"][1]["editable"] is False
    chapter = structure["sections"][3]
    assert chapter["units"][1]["editable"] is True
    if structure["assembled"]:
        # Measured positions in the delivered file, in order, ending with the file.
        clips = [unit["clip"] for unit in chapter["units"]]
        assert all(clip is not None for clip in clips)
        assert all(left[1] <= right[0] + 1e-6 for left, right in zip(clips, clips[1:]))
        audio = client.get(chapter["audio_url"])
        assert audio.status_code == 200 and audio.headers["content-type"] == "audio/mpeg"
    assert client.get(f"/api/jobs/{job_id}/sections/99/audio").status_code == 404


def test_text_from_the_planilla_cannot_be_rewritten_from_the_review(offline) -> None:
    client = TestClient(main.app)
    job_id = client.post("/api/books/jobs", files=_demo_files(),
                         data={"planilla_json": json.dumps(MODEL)}).json()["job_id"]
    asyncio.run(main.run_production_job(job_id))
    units = job_store.get_job(main.database_path, job_id)["sentences"]
    legal = next(unit for unit in units if (unit.get("provenance") or {}).get("role") == "legal")
    refused = client.post(f"/api/jobs/{job_id}/review",
                          json={"sentence_id": legal["sentence_id"], "decision": "correct", "corrected_text": "Otro."})
    assert refused.status_code == 422 and "planilla" in refused.json()["detail"]
    body = next(unit for unit in units if unit["text"].startswith("La barca dejó"))
    accepted = client.post(f"/api/jobs/{job_id}/review",
                           json={"sentence_id": body["sentence_id"], "decision": "correct",
                                 "corrected_text": "La barca dejó a Clara en el muelle."})
    assert accepted.status_code == 200, accepted.text


def test_a_direct_download_has_the_same_structure(offline) -> None:
    client = TestClient(main.app)
    response = client.post("/api/books/package", files=_demo_files(), data={"planilla_json": json.dumps(MODEL)})
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        names = archive.namelist()
        status = json.loads(archive.read("DAISY 3/quality-status.json"))
        assert sum(1 for name in names if re.fullmatch(r"DAISY 3/audio/chapter-\d{3}/chapter-\d{3}\.mp3", name)) == 8
        book = archive.read(next(name for name in names if name.endswith(".xml") and name.count("/") == 1)).decode()
    assert status["state"] == "UNVERIFIED"
    assert {check["status"] for check in status["package"]["once_structure"]} == {"passed"}
    # Headings are read in the ONCE structure (the direct download did not read them before).
    assert 'class="once-title"' in book and "par-heading-0004" in book


def test_without_a_planilla_nothing_changes(offline) -> None:
    client = TestClient(main.app)
    job_id = client.post("/api/books/jobs", files=_demo_files()).json()["job_id"]
    asyncio.run(main.run_production_job(job_id))
    job = job_store.get_job(main.database_path, job_id)
    assert job["state"] == "completed"
    assert "once" not in job["metadata"]
    assert not any((unit.get("provenance") or {}).get("role") for unit in job["sentences"])
    with zipfile.ZipFile(job["output_path"]) as archive:
        book = archive.read(next(name for name in archive.namelist()
                                 if name.endswith(".xml") and name.count("/") == 1)).decode()
    assert "once-" not in book and "Aquí finaliza" not in book
    assert "<pagenum" not in book


def test_the_outline_shows_the_structure_before_anything_is_narrated(offline) -> None:
    client = TestClient(main.app)
    outline = client.post("/api/books/outline", files=_demo_files(), data={"planilla_json": json.dumps(MODEL)}).json()
    assert outline["once"] is True
    assert [section["title"] for section in outline["sections"]][:4] == [
        "El faro de la isla Almendra, De Antonio Pérez", "DATOS BIBLIOGRÁFICOS", "Índice.", "Capítulo 1. La llegada"]
    legal = outline["sections"][0]["units"]
    # Split exactly as the job will narrate it: the heading and the notice's two parts.
    assert [unit["role"] for unit in legal] == ["heading", "legal", "legal"]
    assert " ".join(unit["text"] for unit in legal[1:]) == planilla.ONCE_LEGAL_TEXT
    assert outline["chapters"][0]["title"] == "Capítulo 1. La llegada"
    # A correction made in the editor shows up in the outline.
    corrected = client.post("/api/books/outline", files=_demo_files(), data={
        "planilla_json": json.dumps(MODEL), "corrections_json": json.dumps({"0:0": "Texto corregido."})}).json()
    assert corrected["sections"][3]["units"][1]["text"] == "Texto corregido."
    plain = client.post("/api/books/outline", files=_demo_files()).json()
    assert plain["once"] is False and plain["sections"][0]["title"] == "Capítulo 1. La llegada"


def test_the_jobs_with_different_planillas_are_different_jobs(offline) -> None:
    client = TestClient(main.app)
    first = client.post("/api/books/jobs", files=_demo_files(), data={"planilla_json": json.dumps(MODEL)}).json()
    second = client.post("/api/books/jobs", files=_demo_files(),
                         data={"planilla_json": json.dumps({**MODEL, "index_page": "5"})}).json()
    assert first["job_id"] != second["job_id"] and not second.get("reused_existing_job")


# --------------------------------------------------------------------------- #
# Reading notes, headings and pages from Word
# --------------------------------------------------------------------------- #

def _docx_with_footnote() -> bytes:
    document = Document()
    document.styles.add_style("Page Number DAISY", WD_STYLE_TYPE.CHARACTER)
    document.add_heading("Capítulo 1. La llegada", level=1)
    paragraph = document.add_paragraph("La barca llegó al muelle de madrugada.")
    paragraph.add_run("", style=None)
    document.add_heading("El inventario", level=2)
    marker = document.add_paragraph()
    marker.add_run("12", style="Page Number DAISY")
    document.add_paragraph("Contó tres mantas y dos sillas.")
    buffer = io.BytesIO()
    document.save(buffer)
    # python-docx cannot write footnotes, so the part is added by hand.
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    output = io.BytesIO()
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(output, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.decode("utf-8").replace(
                    "La barca llegó al muelle de madrugada.</w:t></w:r>",
                    "La barca llegó al muelle de madrugada.</w:t></w:r>"
                    '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="2"/></w:r>',
                    1).encode("utf-8")
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.decode("utf-8").replace(
                    "</Relationships>",
                    '<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/footnotes" Target="footnotes.xml"/></Relationships>').encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                data = data.decode("utf-8").replace(
                    "</Types>", '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.'
                    'openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>').encode("utf-8")
            target.writestr(item, data)
        target.writestr("word/footnotes.xml", (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{namespace}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:id="2"><w:p><w:r><w:t>Los faros se encendían a mano.</w:t></w:r></w:p></w:footnote>'
            "</w:footnotes>").encode("utf-8"))
    return output.getvalue()


def test_word_footnotes_headings_and_page_numbers_are_read() -> None:
    structure = extract_book_structure("libro.docx", _docx_with_footnote())
    chapter = structure.chapters[0]
    assert chapter.title == "Capítulo 1. La llegada"
    assert chapter.paragraphs == ["La barca llegó al muelle de madrugada.", "El inventario",
                                  "Contó tres mantas y dos sillas.", "Nota 1. Los faros se encendían a mano."]
    assert chapter.sources[1]["element"] == "heading" and chapter.sources[1]["heading_level"] == 2
    # The page number is a marker where the page begins, never read aloud.
    assert chapter.sources[2]["page_breaks"] == ["12"]
    assert chapter.sources[3]["element"] == "note"
    report = structure.extraction.as_dict()
    assert report["coverage"]["ratio"] == 1.0
    built = once_structure.build(
        [(chapter.title, chapter.paragraphs)], once_structure.selected_provenance([(chapter.title, chapter.paragraphs)],
                                                                                  structure),
        {"title": "Libro", "author": "Autora"}, planilla.Planilla())
    kinds = [(section.kind, section.title) for section in built.sections[3:]]
    assert kinds == [("chapter", "Capítulo 1. La llegada"), ("subchapter", "El inventario")]
    assert built.sections[-1].paragraphs[1].text == "Nota 1. Los faros se encendían a mano. Fin de nota 1."


def test_a_note_that_already_says_nota_is_not_prefixed_twice() -> None:
    source = DEMO.read_text(encoding="utf-8").replace("<p>1. En la isla", "<p>Nota 1. En la isla")
    structure = extract_book_structure("demo.xml", source.encode("utf-8"))
    notes = [text for chapter in structure.chapters for text in chapter.paragraphs if text.startswith("Nota")]
    assert notes == ["Nota 1. En la isla, los faros todavía se encendían a mano cuando fallaba el generador."]
